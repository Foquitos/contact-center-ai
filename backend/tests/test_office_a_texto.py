"""Tests de la conversión de Office a texto (AuditorIA/office_a_texto.py).

POR QUÉ IMPORTA: Gemini acepta los bytes de un .docx/.pptx/.xlsx inline sin dar
error, pero no los ingiere como archivo (medido el 2026-09-03 contra
gemini-3.8-flash: contesta que no ve ningún adjunto, pierde las notas del orador
y con una planilla directamente inventa el contenido; el PDF sí lo lee). Si esta
conversión se rompe, el asistente de docs no falla: arma la documentación sobre
material incompleto o inventado, que es mucho peor.

Los archivos se generan con python-pptx/python-docx/openpyxl SOLO para el test
(el código de producción no depende de ellos: lee el ZIP con la librería
estándar). Si alguno no está instalado, ese test se saltea.
"""
import io
import zipfile

import pytest

from AuditorIA import office_a_texto


def _pptx_de_prueba() -> bytes:
    pptx = pytest.importorskip("pptx", reason="python-pptx es solo para armar el archivo de prueba")
    from pptx.util import Inches

    prs = pptx.Presentation()
    diapo1 = prs.slides.add_slide(prs.slide_layouts[1])
    diapo1.shapes.title.text = "Alta de servicio"
    cuerpo = diapo1.placeholders[1].text_frame
    cuerpo.text = "Paso 1: ingresar al sistema"
    cuerpo.add_paragraph().text = "Paso 2: cargar el DNI"
    diapo1.notes_slide.notes_text_frame.text = "Validar el titular antes de cargar."

    diapo2 = prs.slides.add_slide(prs.slide_layouts[5])
    diapo2.shapes.title.text = "Planes vigentes"
    tabla = diapo2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    tabla.cell(0, 0).text = "Plan"
    tabla.cell(0, 1).text = "Precio"
    tabla.cell(1, 0).text = "Basico"
    tabla.cell(1, 1).text = "$1.000"

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_pptx_trae_titulos_texto_tablas_y_notas():
    texto = office_a_texto.extraer_texto(_pptx_de_prueba(), "instructivo.pptx")

    assert "## Diapositiva 1: Alta de servicio" in texto
    assert "Paso 1: ingresar al sistema" in texto
    assert "Paso 2: cargar el DNI" in texto
    # Las notas del orador suelen tener el detalle que no entra en la diapositiva.
    assert "Validar el titular antes de cargar." in texto
    # La tabla llega como tabla de markdown (el indexador del RAG las preserva).
    assert "| Plan | Precio |" in texto
    assert "| Basico | $1.000 |" in texto


def test_las_diapositivas_salen_en_orden():
    """slide10 va después de slide2, aunque alfabéticamente vaya antes."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        for n in (1, 2, 10):
            z.writestr(f"ppt/slides/slide{n}.xml",
                       f'<p:sld xmlns:p="p" xmlns:a="a"><p:cSld><p:spTree><p:sp><p:txBody>'
                       f'<a:p><a:r><a:t>contenido {n}</a:t></a:r></a:p>'
                       f'</p:txBody></p:sp></p:spTree></p:cSld></p:sld>')

    texto = office_a_texto.extraer_texto(buffer.getvalue(), "orden.pptx")
    assert texto.index("contenido 1") < texto.index("contenido 2") < texto.index("contenido 10")


def test_media_por_diapositiva_ubica_cada_captura():
    """Para que la IA pueda poner la captura en el paso que corresponde, cada
    imagen tiene que saber de qué diapositiva salió."""
    pptx = pytest.importorskip("pptx", reason="python-pptx es solo para armar el archivo de prueba")
    from pptx.util import Inches
    from PIL import Image
    import os

    prs = pptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])          # diapositiva 1: sin imagen
    diapo = prs.slides.add_slide(prs.slide_layouts[6])  # diapositiva 2: con la captura
    imagen = io.BytesIO()
    Image.frombytes("RGB", (300, 300), os.urandom(300 * 300 * 3)).save(imagen, format="PNG")
    imagen.seek(0)
    diapo.shapes.add_picture(imagen, Inches(1), Inches(1), Inches(2), Inches(2))

    buffer = io.BytesIO()
    prs.save(buffer)

    mapa = office_a_texto.media_por_diapositiva(buffer.getvalue())
    assert list(mapa.values()) == [2]
    assert all(ruta.startswith("ppt/media/") for ruta in mapa)


def test_docx_conserva_jerarquia_vinetas_y_tablas():
    docx = pytest.importorskip("docx", reason="python-docx es solo para armar el archivo de prueba")

    documento = docx.Document()
    documento.add_heading("Procedimiento de alta", level=1)
    documento.add_paragraph("Este instructivo explica el alta.")
    documento.add_heading("Requisitos", level=2)
    documento.add_paragraph("DNI vigente", style="List Bullet")
    documento.add_paragraph("Factura de servicio", style="List Bullet")
    tabla = documento.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "Campo"
    tabla.cell(0, 1).text = "Valor"
    tabla.cell(1, 0).text = "Plazo"
    tabla.cell(1, 1).text = "48 horas"

    buffer = io.BytesIO()
    documento.save(buffer)
    texto = office_a_texto.extraer_texto(buffer.getvalue(), "instructivo.docx")

    # La jerarquía de títulos es lo que usa el indexador para armar los fragmentos.
    assert "## Procedimiento de alta" in texto
    assert "### Requisitos" in texto
    assert "- DNI vigente\n- Factura de servicio" in texto
    assert "| Plazo | 48 horas |" in texto


def test_xlsx_sale_como_tabla_de_markdown_por_hoja():
    openpyxl = pytest.importorskip("openpyxl")

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Cartera"
    hoja.append(["Cliente", "Saldo"])
    hoja.append(["ACME", 1500])
    hoja.append([None, None])          # fila vacía: no debería salir
    hoja.append(["Otro", 20])

    buffer = io.BytesIO()
    libro.save(buffer)
    texto = office_a_texto.extraer_texto(buffer.getvalue(), "cartera.xlsx")

    assert "## Hoja: Cartera" in texto
    assert "| Cliente | Saldo |" in texto
    assert "| ACME | 1500 |" in texto
    assert "|  |  |" not in texto


def test_las_barras_de_una_celda_no_rompen_la_tabla():
    assert office_a_texto._celdas_a_markdown([["a|b", "c"], ["1", "2"]]).splitlines()[0] == "| a\\|b | c |"


def test_devuelve_none_para_lo_que_no_es_office():
    """None = 'esto no me toca' (sigue el camino de siempre: inline a Gemini)."""
    assert office_a_texto.extraer_texto(b"%PDF-1.4 contenido", "manual.pdf") is None
    assert office_a_texto.extraer_texto(b"texto suelto", "notas.txt") is None
    assert office_a_texto.extraer_texto(b"", "vacio") is None

    zip_suelto = io.BytesIO()
    with zipfile.ZipFile(zip_suelto, "w") as z:
        z.writestr("hola.txt", "chau")
    assert office_a_texto.extraer_texto(zip_suelto.getvalue(), "cosas.zip") is None


def test_un_office_sin_texto_devuelve_vacio_y_no_none():
    """Una presentación de puras capturas: se leyó y no había texto (las imágenes
    se extraen aparte). No es lo mismo que 'no es un Office'."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("ppt/slides/slide1.xml", '<p:sld xmlns:p="p"><p:cSld><p:spTree/></p:cSld></p:sld>')
    assert office_a_texto.extraer_texto(buffer.getvalue(), "capturas.pptx") == "## Diapositiva 1"


def test_detecta_el_office_viejo_binario():
    assert office_a_texto.es_office_legacy(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"datos")
    assert not office_a_texto.es_office_legacy(b"PK\x03\x04datos")
    assert not office_a_texto.es_office_legacy(b"")


def test_un_archivo_roto_no_tumba_la_corrida():
    """Un ZIP con el XML corrupto devuelve '' (se procesa el resto del material)
    en vez de propagar la excepción."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("word/document.xml", "<w:document> esto no cierra")
    assert office_a_texto.extraer_texto(buffer.getvalue(), "roto.docx") == ""
