"""Convierte documentos Office modernos (.docx / .pptx / .xlsx) a texto plano
para poder mandárselos a Gemini.

POR QUÉ EXISTE
--------------
Gemini acepta los bytes de un .docx/.pptx/.xlsx inline SIN dar error, pero no los
ingiere como archivo. Medido el 2026-09-03 contra gemini-3.8-flash, mandando cada
formato inline y pidiéndole que transcriba todo el texto:

    .pdf   -> lo transcribe completo, sin problemas.
    .pptx  -> contesta que "no se visualiza ningún archivo adjunto" y aun así
              devuelve parte del texto de las diapositivas; las NOTAS DEL ORADOR
              se pierden siempre.
    .docx  -> igual: dice que no hay archivo adjunto y transcribe el cuerpo.
    .xlsx  -> lo peor: no leyó nada de la planilla e INVENTÓ el contenido (con una
              de dos columnas devolvió un plano de un edificio con salas y baños).

Y con la instrucción de contestar "NO_PUEDO" si no podía leer el archivo, los tres
formatos contestaron NO_PUEDO. O sea: el material en Word/PowerPoint/Excel entraba
al asistente de docs y salía un documento incompleto o directamente inventado, sin
ningún error visible. Y no hay a quién delegarle la conversión: la File API (que
del otro lado sí convertía estos formatos) responde 403 con los modelos 3.x, por
eso todo va inline.

Convertido a texto acá, en cambio, llega TODO (notas del orador incluidas), y de
paso sabemos exactamente cuántos caracteres tiene el material, que es lo que
decide si se procesa en una pasada o repartido en varios documentos.

Así que la conversión se hace acá. Los formatos modernos de Office son ZIPs de
XML, con lo cual alcanza la librería estándar (zipfile + ElementTree) y no hace
falta sumar dependencias nuevas al deploy. El Excel usa openpyxl, que ya está en
requirements.txt.

Lo que NO se puede leer es el formato viejo (.doc/.ppt/.xls, binario OLE2): para
eso está `es_office_legacy`, que permite avisarle al usuario que lo guarde como
.docx/.pptx/.xlsx o PDF en vez de dejarlo pasar y procesar aire.

Las imágenes NO salen por acá: las extrae y almacena app.chatbot_imagenes, y a la
IA le llegan como una lista de URLs aparte.
"""
import io
import logging
import re
import zipfile
from typing import Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

logger = logging.getLogger(__name__)

# Firma de un ZIP (todo OOXML lo es) y del Compound File binario de Office 97-2003.
_FIRMA_ZIP = b"PK\x03\x04"
_FIRMA_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# Tope de filas por hoja al convertir una planilla. Una hoja de 50.000 filas no
# entra en el request ni tiene sentido como "documentación": para datos masivos
# está el circuito de tablas de datos (chatbot_tablas_admin), que las carga sin
# pasar por ningún modelo. Se avisa en el texto cuando se recorta, así la
# verificación del asistente lo ve en vez de perder filas en silencio.
MAX_FILAS_POR_HOJA = 5000

_ESTILO_TITULO = re.compile(r"^(?:heading|t[ií]?tulo)\s*[-_]?\s*(\d+)?$", re.IGNORECASE)


# --------------------------------------------------------------------------- #
# Helpers de XML/ZIP                                                           #
# --------------------------------------------------------------------------- #
def _local(tag: str) -> str:
    """Nombre de la etiqueta sin el namespace ('{...}p' -> 'p')."""
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _attr(elem: ET.Element, nombre: str) -> Optional[str]:
    """Atributo por nombre local: los OOXML los prefijan con namespaces distintos
    según el formato (w:val, a:val, r:id...)."""
    for clave, valor in elem.attrib.items():
        if _local(clave) == nombre:
            return valor
    return None


def _abrir(datos: bytes) -> Optional[zipfile.ZipFile]:
    if not datos or not datos.startswith(_FIRMA_ZIP):
        return None
    try:
        return zipfile.ZipFile(io.BytesIO(datos))
    except (zipfile.BadZipFile, OSError, ValueError):
        return None


def _numero_de(nombre: str) -> int:
    """Número que trae el nombre de una parte ('slide10.xml' -> 10). Sirve para
    ordenar: lexicográficamente, slide10 va antes que slide2."""
    m = re.search(r"(\d+)", nombre.rsplit("/", 1)[-1])
    return int(m.group(1)) if m else 0


def _resolver_ruta(carpeta: str, destino: str) -> str:
    """Resuelve el Target relativo de un .rels contra la carpeta de la parte."""
    partes = [p for p in carpeta.split("/") if p]
    for seg in destino.split("/"):
        if seg in ("", "."):
            continue
        if seg == "..":
            if partes:
                partes.pop()
        else:
            partes.append(seg)
    return "/".join(partes)


def _rels_de(z: zipfile.ZipFile, ruta_parte: str) -> List[Tuple[str, str]]:
    """[(tipo, ruta_destino)] de las relaciones internas de una parte."""
    carpeta, nombre = ruta_parte.rsplit("/", 1) if "/" in ruta_parte else ("", ruta_parte)
    try:
        xml = z.read(f"{carpeta}/_rels/{nombre}.rels")
    except KeyError:
        return []
    try:
        raiz = ET.fromstring(xml)
    except ET.ParseError:
        return []

    salida = []
    for rel in raiz:
        destino = rel.get("Target") or ""
        if not destino or rel.get("TargetMode") == "External" or destino.startswith("http"):
            continue
        salida.append((rel.get("Type") or "", _resolver_ruta(carpeta, destino)))
    return salida


def _celdas_a_markdown(filas: List[List[str]]) -> str:
    """Tabla de markdown a partir de una grilla de celdas ya en texto."""
    filas = [f for f in filas if any((c or "").strip() for c in f)]
    if not filas:
        return ""
    ancho = max(len(f) for f in filas)
    norm = [[(c or "").replace("|", "\\|").replace("\n", " ").strip() for c in f] + [""] * (ancho - len(f))
            for f in filas]
    encabezado = norm[0]
    if not any(encabezado):
        encabezado = [f"Columna {i + 1}" for i in range(ancho)]
        cuerpo = norm
    else:
        cuerpo = norm[1:]
    lineas = ["| " + " | ".join(encabezado) + " |",
              "| " + " | ".join(["---"] * ancho) + " |"]
    lineas += ["| " + " | ".join(f) + " |" for f in cuerpo]
    return "\n".join(lineas)


# --------------------------------------------------------------------------- #
# Word (.docx)                                                                 #
# --------------------------------------------------------------------------- #
def _texto_de_parrafo_w(parrafo: ET.Element) -> str:
    """Texto de un <w:p>, respetando tabulaciones y saltos de línea."""
    trozos: List[str] = []
    for nodo in parrafo.iter():
        tag = _local(nodo.tag)
        if tag == "t":
            trozos.append(nodo.text or "")
        elif tag == "tab":
            trozos.append("\t")
        elif tag in ("br", "cr"):
            trozos.append("\n")
    return "".join(trozos).strip()


def _prefijo_de_parrafo_w(parrafo: ET.Element) -> str:
    """'#'*n para los títulos y '- ' para las viñetas/numeradas, según el estilo
    del párrafo: es lo que le da estructura al markdown que después se parte en
    fragmentos para el RAG."""
    ppr = next((h for h in parrafo if _local(h.tag) == "pPr"), None)
    if ppr is None:
        return ""
    estilo = next((_attr(h, "val") or "" for h in ppr if _local(h.tag) == "pStyle"), "")
    m = _ESTILO_TITULO.match(estilo.strip())
    if m:
        nivel = min(int(m.group(1) or 1), 5)
        return "#" * (nivel + 1) + " "
    # La viñeta puede venir en el párrafo (numPr) o heredada del estilo, que es lo
    # que hace Word con "List Bullet" / "Párrafo de lista" (y sus nombres en ES).
    estilo_norm = re.sub(r"[^a-z]", "", estilo.lower())
    if any(_local(h.tag) == "numPr" for h in ppr) or estilo_norm.startswith("list") or "lista" in estilo_norm:
        return "- "
    return ""


def _tabla_w_a_markdown(tabla: ET.Element) -> str:
    filas = []
    for tr in tabla:
        if _local(tr.tag) != "tr":
            continue
        celdas = []
        for tc in tr:
            if _local(tc.tag) != "tc":
                continue
            celdas.append(" ".join(
                _texto_de_parrafo_w(p) for p in tc if _local(p.tag) == "p"
            ).strip())
        filas.append(celdas)
    return _celdas_a_markdown(filas)


def _docx_a_texto(z: zipfile.ZipFile) -> str:
    try:
        raiz = ET.fromstring(z.read("word/document.xml"))
    except (KeyError, ET.ParseError) as e:
        logger.warning("No se pudo leer word/document.xml: %s", e)
        return ""

    cuerpo = next((h for h in raiz if _local(h.tag) == "body"), raiz)
    salida = ""
    venia_lista = False
    for hijo in cuerpo:
        tag = _local(hijo.tag)
        if tag == "p":
            texto = _texto_de_parrafo_w(hijo)
            if not texto:
                continue
            prefijo = _prefijo_de_parrafo_w(hijo)
            es_lista = prefijo == "- "
            # Los ítems seguidos de una lista van pegados: con un renglón en blanco
            # en el medio, el markdown la parte en listas de un solo ítem.
            separador = "" if not salida else ("\n" if es_lista and venia_lista else "\n\n")
            salida += separador + prefijo + texto
            venia_lista = es_lista
        elif tag == "tbl":
            tabla = _tabla_w_a_markdown(hijo)
            if tabla:
                salida += ("\n\n" if salida else "") + tabla
                venia_lista = False
    return salida.strip()


# --------------------------------------------------------------------------- #
# PowerPoint (.pptx)                                                           #
# --------------------------------------------------------------------------- #
def _texto_de_parrafo_a(parrafo: ET.Element) -> str:
    """Texto de un <a:p> (DrawingML: lo usan PowerPoint y las tablas de Word/Excel)."""
    trozos: List[str] = []
    for nodo in parrafo.iter():
        tag = _local(nodo.tag)
        if tag == "t":
            trozos.append(nodo.text or "")
        elif tag == "br":
            trozos.append("\n")
    return "".join(trozos).strip()


def _es_placeholder(forma: ET.Element, tipos: Tuple[str, ...]) -> bool:
    for nodo in forma.iter():
        if _local(nodo.tag) == "ph":
            return (_attr(nodo, "type") or "body") in tipos
    return False


def _tabla_a_markdown(tabla: ET.Element) -> str:
    filas = []
    for tr in tabla:
        if _local(tr.tag) != "tr":
            continue
        celdas = []
        for tc in tr:
            if _local(tc.tag) != "tc":
                continue
            celdas.append(" ".join(
                _texto_de_parrafo_a(p) for p in tc.iter() if _local(p.tag) == "p"
            ).strip())
        filas.append(celdas)
    return _celdas_a_markdown(filas)


def _slides_ordenadas(z: zipfile.ZipFile) -> List[str]:
    return sorted(
        (n for n in z.namelist()
         if n.startswith("ppt/slides/slide") and n.endswith(".xml")),
        key=_numero_de,
    )


def _notas_de_slide(z: zipfile.ZipFile, ruta_slide: str) -> str:
    """Notas del orador de una diapositiva. Se llega por el .rels y no por el
    número: notesSlide3.xml no siempre es la nota de slide3.xml."""
    ruta = next((destino for tipo, destino in _rels_de(z, ruta_slide)
                 if tipo.endswith("/notesSlide")), None)
    if not ruta:
        return ""
    try:
        raiz = ET.fromstring(z.read(ruta))
    except (KeyError, ET.ParseError):
        return ""

    lineas = []
    for forma in raiz.iter():
        if _local(forma.tag) != "sp" or _es_placeholder(forma, ("sldNum", "dt", "ftr")):
            continue
        for parrafo in forma.iter():
            if _local(parrafo.tag) == "p":
                texto = _texto_de_parrafo_a(parrafo)
                if texto:
                    lineas.append(texto)
    return "\n".join(lineas).strip()


def _slide_a_texto(z: zipfile.ZipFile, ruta: str, numero: int) -> str:
    try:
        raiz = ET.fromstring(z.read(ruta))
    except (KeyError, ET.ParseError):
        return ""

    titulo = ""
    bloques: List[str] = []
    for nodo in raiz.iter():
        tag = _local(nodo.tag)
        if tag == "sp":
            lineas = [t for t in (_texto_de_parrafo_a(p) for p in nodo.iter()
                                  if _local(p.tag) == "p") if t]
            if not lineas:
                continue
            if not titulo and _es_placeholder(nodo, ("title", "ctrTitle")):
                titulo = lineas[0]
                lineas = lineas[1:]
            bloques.extend(lineas)
        elif tag == "tbl":
            tabla = _tabla_a_markdown(nodo)
            if tabla:
                bloques.append(tabla)

    encabezado = f"## Diapositiva {numero}" + (f": {titulo}" if titulo else "")
    partes = [encabezado]
    if bloques:
        partes.append("\n".join(bloques))
    notas = _notas_de_slide(z, ruta)
    if notas:
        partes.append(f"Notas del orador:\n{notas}")
    return "\n\n".join(partes).strip()


def _pptx_a_texto(z: zipfile.ZipFile) -> str:
    slides = _slides_ordenadas(z)
    if not slides:
        return ""
    textos = [_slide_a_texto(z, ruta, i) for i, ruta in enumerate(slides, start=1)]
    return "\n\n".join(t for t in textos if t).strip()


def media_por_diapositiva(datos: bytes) -> Dict[str, int]:
    """{'ppt/media/image3.png': 4} — en qué diapositiva aparece cada imagen.

    Con esto, la captura que se guarda en la base lleva el número de diapositiva
    en el nombre y la IA puede ubicarla en el paso que corresponde, en vez de
    recibir una lista de 'image7.png' sin contexto. Si una imagen aparece en
    varias diapositivas se queda con la primera.
    """
    z = _abrir(datos)
    if z is None:
        return {}
    mapa: Dict[str, int] = {}
    with z:
        for numero, ruta in enumerate(_slides_ordenadas(z), start=1):
            for tipo, destino in _rels_de(z, ruta):
                if tipo.endswith("/image") and destino not in mapa:
                    mapa[destino] = numero
    return mapa


# --------------------------------------------------------------------------- #
# Excel (.xlsx)                                                                #
# --------------------------------------------------------------------------- #
def _xlsx_a_texto(datos: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - openpyxl está en requirements.txt
        logger.warning("openpyxl no está instalado: no se puede convertir la planilla.")
        return ""

    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as e:
        logger.warning("No se pudo abrir la planilla: %s", e)
        return ""

    bloques: List[str] = []
    try:
        for hoja in libro.worksheets:
            filas: List[List[str]] = []
            recortada = False
            for i, fila in enumerate(hoja.iter_rows(values_only=True)):
                if i >= MAX_FILAS_POR_HOJA:
                    recortada = True
                    break
                filas.append(["" if c is None else str(c).strip() for c in fila])
            tabla = _celdas_a_markdown(filas)
            if not tabla:
                continue
            bloque = f"## Hoja: {hoja.title}\n\n{tabla}"
            if recortada:
                bloque += (f"\n\n> ATENCIÓN: la hoja tiene más de {MAX_FILAS_POR_HOJA} filas y "
                           f"acá solo están las primeras {MAX_FILAS_POR_HOJA}.")
            bloques.append(bloque)
    finally:
        libro.close()
    return "\n\n".join(bloques).strip()


# --------------------------------------------------------------------------- #
# API pública                                                                  #
# --------------------------------------------------------------------------- #
def es_office_legacy(datos: Optional[bytes]) -> bool:
    """True si es un Office binario viejo (.doc/.xls/.ppt de Office 97-2003), que
    no se puede leer ni acá ni del lado de Gemini."""
    return bool(datos) and datos.startswith(_FIRMA_OLE2)


def extraer_texto(datos: Optional[bytes], nombre: str = "") -> Optional[str]:
    """Texto de un .docx/.pptx/.xlsx, o None si el archivo no es ninguno de esos.

    Devuelve '' (y no None) cuando SÍ es un Office moderno pero no trae texto:
    por ejemplo una presentación de puras capturas. Quien llama distingue así
    'esto no me toca' de 'lo leí y no había texto'.
    """
    z = _abrir(datos)
    if z is None:
        return None

    with z:
        partes = set(z.namelist())
        try:
            if "word/document.xml" in partes:
                return _docx_a_texto(z)
            if any(p.startswith("ppt/slides/slide") for p in partes):
                return _pptx_a_texto(z)
            if "xl/workbook.xml" in partes:
                return _xlsx_a_texto(datos)
        except Exception:
            logger.exception("Error convirtiendo el archivo Office '%s' a texto.", nombre or "?")
            return ""
    return None  # Es un ZIP, pero no un Office (un .zip suelto, un .apk, etc.)
