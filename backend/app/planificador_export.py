"""Planificador: los mismos números de la pantalla, en un Excel.

POR QUÉ ESTO EXISTE
-------------------
La planificación no se consume sola: se manda, se discute con el cliente, se
pega en una reunión y se cruza contra la malla en la planilla de siempre. Una
pantalla no sirve para nada de eso. Y como Voltara se dimensionó durante años en
Excel, poder abrir el resultado ahí es lo que permite comparar el número nuevo
contra el viejo en el mismo lugar.

DOS DECISIONES QUE NO SON OBVIAS
--------------------------------
1. **Va siempre una hoja de parámetros.** Un Excel con la dotación y sin los
   supuestos es un número sin defensa: dentro de dos semanas nadie se va a
   acordar con qué shrinkage ni con qué paciencia salió. La hoja "Parámetros"
   viaja con el archivo para que la planificación se pueda auditar sin volver al
   sistema.
2. **Los porcentajes van como número con formato de porcentaje, no como texto.**
   Un "80,3%" escrito como texto no se puede promediar ni graficar, y el archivo
   existe justamente para que alguien lo siga trabajando.

No usa pandas: son planillas chicas y openpyxl da control sobre el formato, que
acá es la mitad del entregable.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, replace
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MEDIA_XLSX = ("application/vnd.openxmlformats-officedocument"
              ".spreadsheetml.sheet")

# Formatos de celda. El de porcentaje con un decimal porque un NDS de 80,3% y
# uno de 80% no son lo mismo y redondear los confunde.
PCT = "0.0%"
PCT2 = "0.00%"
NUM0 = "#,##0"
NUM1 = "#,##0.0"
NUM2 = "#,##0.00"
FECHA = "dd/mm/yyyy"
FECHA_HORA = "dd/mm/yyyy hh:mm"

_CABECERA = PatternFill("solid", fgColor="EFEFEB")
_TITULO = Font(bold=True, size=12)
_NEGRITA = Font(bold=True)
_TENUE = Font(color="6A6A66", italic=True)


@dataclass
class Columna:
    clave: str
    titulo: str
    formato: Optional[str] = None
    ancho: int = 14


def _escribir_tabla(ws, columnas: Sequence[Columna], filas: Sequence[dict],
                    fila_inicial: int = 1) -> int:
    """Escribe una tabla con encabezado fijo. Devuelve la fila siguiente."""
    for i, col in enumerate(columnas, start=1):
        celda = ws.cell(row=fila_inicial, column=i, value=col.titulo)
        celda.font = _NEGRITA
        celda.fill = _CABECERA
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = col.ancho

    for j, fila in enumerate(filas, start=fila_inicial + 1):
        for i, col in enumerate(columnas, start=1):
            celda = ws.cell(row=j, column=i, value=_valor(fila.get(col.clave)))
            if col.formato and celda.value is not None:
                celda.number_format = col.formato

    ws.freeze_panes = ws.cell(row=fila_inicial + 1, column=1)
    ws.auto_filter.ref = (f"A{fila_inicial}:"
                          f"{get_column_letter(len(columnas))}"
                          f"{fila_inicial + max(len(filas), 1)}")
    return fila_inicial + len(filas) + 2


def _valor(v: Any):
    """openpyxl no sabe escribir Decimal ni tipos raros; los números tienen que
    llegar como números o el formato de porcentaje no se aplica."""
    if v is None or isinstance(v, (int, float, str, datetime, date)):
        return v
    try:
        return float(v)
    except (TypeError, ValueError):
        return str(v)


def _bloque(ws, fila: int, titulo: str, pares: Sequence[tuple],
            nota: Optional[str] = None) -> int:
    """Un bloque de clave/valor con título. Para las hojas de resumen."""
    c = ws.cell(row=fila, column=1, value=titulo)
    c.font = _TITULO
    fila += 1
    for etiqueta, valor, *formato in pares:
        ws.cell(row=fila, column=1, value=etiqueta).font = _NEGRITA
        celda = ws.cell(row=fila, column=2, value=_valor(valor))
        if formato and formato[0] and celda.value is not None:
            celda.number_format = formato[0]
        fila += 1
    if nota:
        celda = ws.cell(row=fila, column=1, value=nota)
        celda.font = _TENUE
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 2, end_column=6)
        fila += 3
    return fila + 1


def _guardar(wb: Workbook) -> io.BytesIO:
    flujo = io.BytesIO()
    wb.save(flujo)
    flujo.seek(0)
    return flujo


# =========================================================================
# LIBRO DEL PLAN
# =========================================================================

_COLS_REQ = [
    Columna("dia", "Día", FECHA, 12),
    Columna("hora", "Intervalo", None, 10),
    Columna("Pool", "Pool", None, 24),
    Columna("Llamadas", "Llamadas", NUM1, 11),
    Columna("TmoSeg", "TMO (seg)", NUM0, 11),
    Columna("Trafico", "Erlangs", NUM2, 10),
    Columna("OperadoresLinea", "En línea", NUM0, 10),
    Columna("OperadoresPlanificar", "A planificar", NUM0, 12),
    Columna("OperadoresPlanificados", "Citados", NUM0, 10),
    Columna("CitadosEquivalentes", "Citados equivalentes (antigüedad)", NUM1, 20),
    Columna("Brecha", "Brecha", NUM0, 10),
    Columna("RefuerzoDisponible", "Digital con turno", NUM0, 17),
    Columna("RefuerzoCubre", "Lo cubre Digital", NUM0, 16),
    Columna("FaltanteNeto", "Falta tras Digital", NUM0, 17),
    Columna("NdsContractual", "NDS", PCT, 9),
    Columna("AsaSeg", "ASA (seg)", NUM1, 11),
    Columna("Abandono", "Abandono", PCT2, 11),
    Columna("NivelAtencionB", "Niv. atención (Erlang B)", PCT, 20),
    Columna("Ocupacion", "Ocupación", PCT, 11),
    Columna("Motivo", "Manda", None, 26),
]

_COLS_PRON = [
    Columna("dia", "Día", FECHA, 12),
    Columna("hora", "Intervalo", None, 10),
    Columna("skill", "Skill", None, 26),
    Columna("LlamadasTotal", "Demanda del cliente", NUM1, 18),
    Columna("Asignacion", "Nos asignan", PCT, 12),
    Columna("LlamadasBase", "Nuestras (base)", NUM1, 15),
    Columna("LlamadasAcme", "Nuestras (con ajustes)", NUM1, 20),
    Columna("Cliente", "Pronóstico del cliente", NUM1, 20),
    Columna("Real", "Real (lo que entró)", NUM1, 18),
    Columna("TmoSeg", "TMO (seg)", NUM0, 11),
]


def _con_refuerzo(columnas: Sequence[Columna], etiqueta: str) -> List[Columna]:
    """Las columnas con el nombre del refuerzo de la campaña en vez de «Digital»,
    que es como se llama en Voltara."""
    if etiqueta == "Digital":
        return list(columnas)
    return [replace(c, titulo=c.titulo.replace("Digital", etiqueta)) for c in columnas]


def libro_del_plan(plan: Dict[str, Any], config: Dict[str, Any],
                   campana: str = "") -> io.BytesIO:
    """Requerimiento, pronóstico y los supuestos con los que salieron."""
    ref = (config.get("refuerzo") or {}).get("etiqueta") or "Digital"
    wb = Workbook()
    corrida = plan.get("corrida") or {}
    metricas = corrida.get("Metricas") or {}

    ws = wb.active
    ws.title = "Resumen"
    fila = _bloque(ws, 1, f"Planificación{' — ' + campana if campana else ''}", [
        ("Corrida", corrida.get("CorridaID"), NUM0),
        ("Horizonte", corrida.get("Horizonte")),
        ("Desde", corrida.get("Desde"), FECHA),
        ("Hasta", corrida.get("Hasta"), FECHA),
        ("Modelo de volumen", corrida.get("Modelo")),
        ("Calculada el", corrida.get("TerminadoEn") or corrida.get("CreadoEn"), FECHA_HORA),
        ("Pronóstico sobre la demanda total del cliente",
         "sí" if plan.get("sobre_demanda_total") else "no"),
    ])

    brecha = metricas.get("brecha") or {}
    fila = _bloque(ws, fila, "Cifras del período", [
        ("Intervalos", metricas.get("intervalos"), NUM0),
        ("Llamadas pronosticadas", metricas.get("llamadas"), NUM0),
        ("Pico de operadores a planificar", metricas.get("pico_operadores"), NUM0),
        ("Horas-operador", metricas.get("horas_operador"), NUM1),
        ("NDS promedio proyectado", metricas.get("nds_promedio"), PCT),
        ("Abandono promedio proyectado", metricas.get("abandono_promedio"), PCT2),
    ], nota=(
        "«En línea» es lo que pide la cola. «A planificar» ya tiene aplicadas la "
        "disponibilidad medida y el shrinkage de nómina, en ese orden y nunca "
        "sumados: son dos descuentos distintos y combinarlos descuenta dos veces "
        "lo mismo."))

    if brecha:
        fila = _bloque(ws, fila, "Contra la malla de RRHH", [
            ("Intervalos con malla cargada", brecha.get("intervalos_con_malla"), NUM0),
            ("Intervalos en falta", brecha.get("intervalos_en_falta"), NUM0),
            ("Intervalos con exceso", brecha.get("intervalos_con_exceso"), NUM0),
            ("Operadores faltantes en el pico", brecha.get("operadores_faltantes_pico"), NUM0),
            ("Horas-operador faltantes", brecha.get("horas_operador_faltantes"), NUM1),
            ("Horas-operador sobrantes", brecha.get("horas_operador_sobrantes"), NUM1),
        ], nota=(
            "Las dos puntas van por separado y no netas: un día con dos horas de "
            "faltante y dos de sobrante no está equilibrado, está mal armado, y el "
            "neto en cero lo escondería."))

    if brecha.get("con_refuerzo"):
        fila = _bloque(ws, fila, f"Si no se llega: pasar gente de {ref}", [
            (f"Horas-operador que cubre {ref}",
             brecha.get("horas_operador_cubre_refuerzo"), NUM1),
            ("Horas-operador que faltan igual",
             brecha.get("horas_operador_faltantes_netas"), NUM1),
            ("Intervalos que no se cubren ni así", brecha.get("intervalos_sin_cubrir"), NUM0),
            (f"Faltan en el peor intervalo, tras {ref}",
             brecha.get("operadores_faltantes_netos_pico"), NUM0),
        ], nota=(
            f"{ref} es back office sin SLA: su gente se puede pasar al teléfono "
            f"cuando la malla no alcanza. «Lo cubre {ref}» es lo menor entre lo que "
            f"falta y la gente de {ref} con turno en ese intervalo. No cambia «A "
            "planificar»: lo que pide la cola es lo mismo, cambia quién lo cubre. Lo "
            "que falta igual es lo que obliga a mover la malla o pedir horas extra."))

    avisos = list(metricas.get("avisos") or [])
    if avisos:
        ws.cell(row=fila, column=1, value="Avisos de la corrida").font = _TITULO
        fila += 1
        for a in avisos:
            celda = ws.cell(row=fila, column=1, value=f"• {a}")
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=8)
            fila += 1

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 24

    nombres = {s["skill_id"]: s["nombre"]
               for p in (config.get("pools") or []) for s in (p.get("skills") or [])}

    _escribir_tabla(wb.create_sheet("Requerimiento"), _con_refuerzo(_COLS_REQ, ref),
                    [_con_fecha(r) for r in plan.get("requerimiento") or []])
    # Las mismas tres series que la pantalla, en la misma fila: quien baja el
    # Excel para discutir una dotación necesita poder ver al lado lo que
    # pronosticó el cliente y lo que terminó entrando, o vuelve a la pantalla a
    # copiarlo a mano.
    cliente = _indice_de_cotejo(plan.get("cliente"))
    real = _indice_de_cotejo(plan.get("real"))
    hasta = plan.get("real_hasta")
    _escribir_tabla(wb.create_sheet("Pronóstico"), _COLS_PRON,
                    [_con_fecha(_con_cotejo(f, cliente, real, hasta), nombres)
                     for f in plan.get("pronostico") or []])
    _hoja_parametros(wb.create_sheet("Parámetros"), config)
    return _guardar(wb)


def _indice_de_cotejo(filas) -> Dict[tuple, float]:
    return {(f["Intervalo"], f["SkillID"]): f["Llamadas"] for f in filas or []}


def _con_cotejo(fila: dict, cliente: Dict[tuple, float],
                real: Dict[tuple, float], hasta) -> dict:
    """Le pega a la fila del pronóstico lo que pronosticó el cliente y lo que
    entró, si están.

    El real queda VACÍO —no en cero— donde todavía no hay dato: en una tabla
    dinámica un cero se suma igual que un dato y hunde el total del día.
    """
    clave = (fila.get("Intervalo"), fila.get("SkillID"))
    salida = dict(fila)
    salida["Cliente"] = cliente.get(clave)
    salida["Real"] = (real.get(clave, 0.0)
                      if hasta is not None and clave[0] is not None
                      and clave[0] <= hasta else None)
    return salida


def _con_fecha(fila: dict, nombres: Optional[Dict[int, str]] = None) -> dict:
    """Parte el datetime en día y hora: en Excel se filtra por día y se ordena
    por hora, y un solo campo obliga a hacer las dos cosas mal."""
    salida = dict(fila)
    momento = fila.get("Intervalo") or fila.get("momento")
    if isinstance(momento, datetime):
        salida["dia"] = momento.date()
        salida["hora"] = momento.strftime("%H:%M")
    if nombres is not None:
        salida["skill"] = nombres.get(fila.get("SkillID"), fila.get("SkillID"))
    return salida


_COLS_SKILL = [
    Columna("pool", "Pool", None, 24),
    Columna("nombre", "Skill", None, 26),
    Columna("objetivo_nds", "Objetivo NDS", PCT, 13),
    Columna("umbral_seg", "Umbral (seg)", NUM0, 12),
    Columna("max_abandono", "Techo de abandono", PCT, 17),
    Columna("paciencia_seg", "Paciencia (seg)", NUM0, 14),
    Columna("max_asa_seg", "TME máximo (seg)", NUM0, 16),
    Columna("objetivo_nds_2", "2º NDS", PCT, 10),
    Columna("umbral_seg_2", "2º umbral (seg)", NUM0, 14),
    Columna("min_nivel_atencion_b", "Nivel de atención mínimo", PCT, 22),
    Columna("llamadas_30d", "Llamadas últimos 30 días", NUM0, 22),
]


def _hoja_parametros(ws, config: Dict[str, Any]) -> None:
    proc = config.get("procedencia") or {}
    fila = _bloque(ws, 1, "Parámetros con los que se calculó", [
        ("Intervalo (minutos)", config.get("intervalo_min"), NUM0),
        ("Techo de ocupación", config.get("max_ocupacion"), PCT),
        ("Shrinkage de nómina", config.get("shrinkage"), PCT),
        ("  — de ausentismo", proc.get("shrinkage_ausentismo"), PCT),
        ("  — de capacitación", proc.get("shrinkage_capacitacion"), PCT),
        ("Origen del shrinkage", proc.get("shrinkage_origen")),
        ("Medido el", proc.get("shrinkage_medido_en"), FECHA),
        ("Paciencia del cliente (seg)", config.get("paciencia_seg"), NUM0),
        ("Origen de la paciencia", proc.get("paciencia_origen")),
        ("Horizonte de ajuste de la paciencia (seg)",
         proc.get("paciencia_horizonte_seg"), NUM0),
    ], nota=(
        "La paciencia no es «cuánto espera un cliente»: es el parámetro de la "
        "exponencial que reproduce el abandono observado en los primeros segundos, "
        "que es donde se juega el dimensionamiento. Por eso el número es más grande "
        "que lo que cualquiera esperaría a mano."))

    filas = []
    for p in config.get("pools") or []:
        for s in p.get("skills") or []:
            filas.append({**s, "pool": p.get("nombre")})
    fila = _escribir_tabla(ws, _COLS_SKILL, filas, fila_inicial=fila)

    franjas = config.get("disponibilidad") or []
    if franjas:
        dias = ["Todos", "Lunes", "Martes", "Miércoles", "Jueves", "Viernes",
                "Sábado", "Domingo"]
        ws.cell(row=fila, column=1, value="Disponibilidad por franja").font = _TITULO
        fila = _escribir_tabla(ws, [
            Columna("dia", "Día", None, 14),
            Columna("hora_desde", "Desde (hora)", NUM0, 12),
            Columna("hora_hasta", "Hasta (hora)", NUM0, 12),
            Columna("factor", "Disponibilidad", PCT, 14),
        ], [{**f, "dia": dias[f.get("dia_semana", 0)]} for f in franjas],
            fila_inicial=fila + 1)


# =========================================================================
# LIBRO DEL BACKTEST
# =========================================================================

_COLS_DIA = [
    Columna("dia", "Día", FECHA, 12),
    Columna("dia_semana", "Día de semana", None, 14),
    Columna("real", "Real (nuestras)", NUM0, 15),
    Columna("pronosticado", "Pronosticado", NUM0, 14),
    Columna("diferencia", "Diferencia", NUM0, 12),
    Columna("desvio", "Desvío", PCT, 10),
    Columna("ingenuo", "Semana anterior", NUM0, 15),
    Columna("produccion", "dbo.Forecast", NUM0, 13),
    Columna("desvio_produccion", "Desvío dbo.Forecast", PCT, 19),
    Columna("combinado", "Combinado", NUM0, 12),
    Columna("peso_cliente", "Factor de la mezcla", NUM2, 19),
    Columna("tipo_dia", "Tipo de día", None, 12),
    Columna("factor_clima", "Factor de clima", NUM2, 15),
    Columna("real_total", "Demanda real del cliente", NUM0, 22),
    Columna("pronosticado_total", "Demanda pronosticada", NUM0, 20),
    Columna("share_real", "Share real", PCT, 11),
    Columna("share_aplicado", "Share aplicado", PCT, 14),
    Columna("evento", "Evento", None, 22),
    Columna("parcial", "Día incompleto", None, 14),
]

_COLS_INTERVALO = [
    Columna("dia", "Día", FECHA, 12),
    Columna("hora", "Intervalo", None, 10),
    Columna("real", "Real (nuestras)", NUM1, 15),
    Columna("pronosticado", "Pronosticado", NUM1, 14),
    Columna("diferencia", "Diferencia", NUM1, 12),
    Columna("produccion", "dbo.Forecast", NUM1, 13),
    Columna("combinado", "Combinado", NUM1, 12),
    Columna("real_total", "Demanda real del cliente", NUM1, 22),
    Columna("pronosticado_total", "Demanda pronosticada", NUM1, 20),
]

_ETIQUETA_RASGO = {
    "cdd": "calor (grados sobre 24 aparentes)",
    "cdd2": "calor, al cuadrado",
    "hdd": "frío (grados bajo 14 aparentes)",
    "hdd2": "frío, al cuadrado",
    "cdd_ayer": "calor del día anterior",
    "hdd_ayer": "frío del día anterior",
    "lluvia": "lluvia (log de los mm)",
    "rafaga": "ráfagas sobre 40 km/h",
    "cdd_3d": "calor acumulado de 3 días",
    "hdd_3d": "frío acumulado de 3 días",
    "cdd_7d": "calor acumulado de 7 días",
    "hdd_7d": "frío acumulado de 7 días",
    "humedad": "humedad (sobre 60%)",
    "viento": "viento sostenido",
    "amplitud": "amplitud térmica del día",
    "anual_sin": "estacionalidad anual (seno)",
    "anual_cos": "estacionalidad anual (coseno)",
    "anual_sin2": "estacionalidad semestral (seno)",
    "anual_cos2": "estacionalidad semestral (coseno)",
    "dia_mes_sin": "ciclo de facturación (seno)",
    "dia_mes_cos": "ciclo de facturación (coseno)",
    "vispera": "víspera de feriado",
    "post_feriado": "día después de un feriado",
    "finde_largo": "fin de semana largo",
    "vacaciones": "enero o vacaciones de invierno",
}

_DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


def libro_del_backtest(bt: Dict[str, Any], campana: str = "",
                       refuerzo: str = "Digital") -> io.BytesIO:
    """Lo pronosticado contra lo que pasó, día por día e intervalo por intervalo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    res = bt.get("resumen") or {}
    intervalo = res.get("intervalo") or {}
    diario = res.get("diario") or {}
    sin_ev = res.get("diario_sin_eventos") or {}
    forma = res.get("forma") or {}
    total = res.get("total_cliente") or {}
    ingenuo = res.get("ingenuo") or {}
    prod = res.get("produccion") or {}
    prod_dia = res.get("produccion_diario") or {}

    fila = _bloque(ws, 1, f"Backtest del pronóstico{' — ' + campana if campana else ''}", [
        ("Desde", bt.get("desde"), FECHA),
        ("Hasta", bt.get("hasta"), FECHA),
        ("Antelación (días)", bt.get("antelacion"), NUM0),
        ("Días cerrados evaluados", bt.get("dias_evaluados"), NUM0),
        ("Modelo", bt.get("modelo")),
    ], nota=(
        "Cada día se pronosticó con la historia que había «antelación» días antes. "
        "No se aplicaron los ajustes manuales del pronóstico: acá se mide el modelo, "
        "no la corrección de una persona. Nada de esto se guardó en la base."))

    fila = _bloque(ws, fila, "Error sobre NUESTRAS llamadas", [
        ("Real", diario.get("real"), NUM0),
        ("Pronosticado", diario.get("pronosticado"), NUM0),
        ("Sesgo (positivo = el pronóstico se quedó corto)", diario.get("sesgo"), PCT),
        ("WAPE por intervalo", intervalo.get("wape"), PCT),
        ("MAPE diario", diario.get("mape"), PCT),
        ("MAPE diario sin días con evento", sin_ev.get("mape"), PCT),
        ("WAPE de la CURVA, sabiendo el total del día", forma.get("wape"), PCT),
        ("WAPE de repetir la semana anterior", ingenuo.get("wape"), PCT),
        ("Mejora contra repetir la semana anterior",
         res.get("mejora_vs_ingenuo"), PCT),
        ("WAPE de dbo.Forecast (el que usa la operación)", prod.get("wape"), PCT),
        ("MAPE diario de dbo.Forecast", prod_dia.get("mape"), PCT),
        ("Mejora contra dbo.Forecast", res.get("mejora_vs_produccion"), PCT),
    ], nota=(
        "WAPE = suma de los errores absolutos sobre el total real: pondera por "
        "volumen, así que dice cuánto se erró donde hay gente atendiendo. MAPE = "
        "promedio de los errores relativos: sobre medias horas de madrugada explota "
        "y no significa nada, por eso se informa sobre totales diarios. La mejora "
        "contra la semana anterior es la que dice si el modelo aporta algo: si da "
        "cero, no está haciendo nada que no haga repetir el mismo día de la semana pasada. "
        "El «WAPE de la CURVA» es lo que quedaría de error si el total del día se "
        "supiera de antemano: separa fallar en CUÁNTAS llamadas entran —que se ataca "
        "con clima, feriados y eventos— de fallar en CUÁNDO entran, que se ataca con el "
        "perfil de horas y es lo que rompe la malla de turnos. dbo.Forecast es el "
        "pronóstico que la operación ya está usando: si el planificador no le gana, no "
        "hay motivo para cambiar."))

    if total:
        fila = _bloque(ws, fila, "Error sobre la demanda TOTAL del cliente", [
            ("Real", total.get("real"), NUM0),
            ("Pronosticado", total.get("pronosticado"), NUM0),
            ("Sesgo", total.get("sesgo"), PCT),
            ("MAPE diario", total.get("mape"), PCT),
        ], nota=(
            "Este es el error del modelo propiamente dicho. El de arriba le suma el "
            "reparto —qué porcentaje de las llamadas del cliente nos toca—, que sale "
            "de tramos cargados mirando la historia. Si este número es bueno y el de "
            "arriba no, el problema no es el pronóstico: es el share."))

    cl = bt.get("clima") or {}
    if cl.get("usado"):
        m = cl.get("modelo") or {}
        co = m.get("coeficientes") or {}
        fila = _bloque(ws, fila, "Modelo de clima", [
            ("Días observados con los que se entrenó", m.get("n"), NUM0),
            ("Entrenado hasta", m.get("entrenado_hasta"), FECHA),
            ("Variación del nivel diario que explica (R2)", m.get("r2"), PCT),
            ("Cuánto baja el desvío del residuo", m.get("reduccion_residuo"), PCT),
            ("Factor mínimo aplicado", cl.get("factor_min"), NUM2),
            ("Factor máximo aplicado", cl.get("factor_max"), NUM2),
        ] + [(f"  efecto por unidad — {_ETIQUETA_RASGO.get(k, k)}",
              math.exp(v) - 1, PCT2)
             for k, v in co.items() if k != "intercepto"],
            nota=(
            "La demanda de una distribuidora eléctrica sigue una U, no una recta: los "
            "días de mayor volumen son los de mucho calor Y los de mucho frío. Por eso "
            "hay grados-día de refrigeración y de calefacción por separado. El factor "
            "multiplica el perfil estacional ANTES de la corrección de nivel, y la "
            "corrección se calcula sobre lo que el clima no explicó."))
    elif cl:
        fila = _bloque(ws, fila, "Modelo de clima", [
            ("Usado", "no"), ("Motivo", cl.get("motivo")),
        ])

    dg = ((bt.get("dotacion") or {}).get("resumen")) or {}
    if dg:
        dd = bt.get("dotacion") or {}
        fila = _bloque(ws, fila, "Operadores: cuántos pedíamos y cuántos hubo", [
            ("Días con volumen", dg.get("dias"), NUM0),
            ("EN LÍNEA — pedíamos (con el pronóstico)", dg.get("en_linea_pron"), NUM0),
            ("EN LÍNEA — hacía falta (con lo que entró)", dg.get("en_linea_real"), NUM0),
            ("EN LÍNEA — pedíamos vs hacía falta", dg.get("linea_pron_vs_real"), PCT),
            ("Conectados logueados (total, con pausas)", dg.get("conectados"), NUM0),
            ("Hacía falta en línea vs logueados",
             dg.get("linea_real_vs_conectados"), PCT),
            ("A CITAR — pedíamos (con el pronóstico)", dg.get("a_planificar_pron"), NUM0),
            ("A CITAR — hacía falta (con lo que entró)", dg.get("a_planificar_real"), NUM0),
            ("A CITAR — tenían turno (registro del día cerrado, con extras)",
             dg.get("citados"), NUM0),
            ("A CITAR — malla (lo publicado antes del día, sin extras)", dg.get("malla"), NUM0),
            ("A CITAR — turnos de más que la malla", dg.get("registro_vs_malla"), PCT),
            ("Conectados vs los que tenían turno",
             dg.get("conectados_vs_citados"), PCT),
            ("  telefónicos (con Contingencia, SVP y Anfitrión cuando atienden)",
             dg.get("conectados_pool"), NUM0),
            (f"  de {refuerzo}", dg.get("conectados_refuerzo"), NUM0),
            ("  de otras digitales", dg.get("conectados_digital"), NUM0),
            ("  de colas dedicadas (T1 - Consumo)", dg.get("conectados_dedicada"), NUM0),
            ("  sin clasificar", dg.get("conectados_otras"), NUM0),
            ("EN LÍNEA, SIN PAUSAS — telefónicos", dg.get("conectados_pool_en_linea"), NUM0),
            (f"EN LÍNEA, SIN PAUSAS — {refuerzo}", dg.get("conectados_refuerzo_en_linea"), NUM0),
            ("EN LÍNEA, SIN PAUSAS — otras digitales",
             dg.get("conectados_digital_en_linea"), NUM0),
            ("EN LÍNEA, SIN PAUSAS — colas dedicadas",
             dg.get("conectados_dedicada_en_linea"), NUM0),
            ("EN LÍNEA, SIN PAUSAS — sin clasificar",
             dg.get("conectados_otras_en_linea"), NUM0),
            ("Tenían turno: de Contingencia, SVP y Anfitrión en la línea",
             dg.get("citados_parciales"), NUM0),
            (f"{refuerzo} con turno", dg.get("refuerzo_turno"), NUM0),
            (f"De {refuerzo}, parte en la línea", dg.get("refuerzo_en_linea"), PCT),
            ("Faltó (sumado intervalo por intervalo)", dg.get("faltante"), NUM0),
            (f"  lo cubría {refuerzo}", dg.get("cubre_refuerzo"), NUM0),
            ("  faltó igual", dg.get("faltante_neto"), NUM0),
            ("A CITAR — pedíamos vs tenían turno", dg.get("pron_vs_citados"), PCT),
            ("A CITAR — hacía falta vs tenían turno", dg.get("real_vs_citados"), PCT),
            ("A CITAR — de más por el pronóstico", dg.get("de_mas_por_pronostico"), NUM0),
            ("Nivel de servicio logrado", dg.get("nds_real"), PCT),
            ("Objetivo de nivel de servicio", dd.get("objetivo_nds"), PCT),
            ("Abandono real", dg.get("abandono_real"), PCT2),
            ("CPH pronosticado (llamadas pronosticadas / horas pedidas)",
             dg.get("cph_pron"), NUM2),
            ("CPH real (llamadas atendidas / horas conectadas)",
             dg.get("cph_real"), NUM2),
            ("Ocupación del modelo", dg.get("ocupacion_modelo"), PCT),
            ("Ocupación real", dg.get("ocupacion_real"), PCT),
            ("Techo de ocupación configurado", dg.get("max_ocupacion"), PCT),
            ("Del pedido, lo fija el techo de ocupación",
             dg.get("peso_techo_ocupacion"), PCT),
        ], nota=(
            "Hay DOS UNIDADES y no se mezclan. EN LÍNEA son operadores sobre la cola: "
            "pedíamos, hacía falta y conectados, comparables entre sí. A CITAR son "
            "personas con turno, que ya llevan adentro disponibilidad, shrinkage y "
            "break: pedíamos, hacía falta, tenían turno y malla, comparables entre sí. "
            "Un número en línea contra uno a citar siempre da una diferencia que no "
            "existe. "
            "Todo en intervalos-operador: la suma de los 48 intervalos del día, que "
            "es la unidad en la que se compran horas. Cada intervalo se dimensiona DOS "
            "veces —con el pronóstico de ese día y con la demanda que realmente "
            "entró— porque el error de llamadas no se traduce uno a uno en error de "
            "gente: Erlang no es lineal y la cadena de descuentos agrega su propio "
            "error. Así la brecha contra la malla se parte en error de PRONÓSTICO "
            "(pedíamos − hacía falta) y error de DIMENSIONAMIENTO (hacía falta − "
            "citados), que se arreglan en lugares distintos. «Tenían turno» sale del "
            "REGISTRO (dbo.payroll) y no de la malla: payroll_futuro no tiene ninguna "
            "hora extra cargada —90 días del pool telefónico dan 1.426,4 h de extras "
            "en el registro y cero en la malla—. Neto son +1,0% de intervalos, pero "
            "día por día llega a +24,6% en un domingo de refuerzo. El nivel de servicio va "
            "al lado porque sin él la brecha se lee al revés: citar menos de lo que el "
            "plan pedía y cumplir el objetivo igual significa que el plan pide de más. "
            "Y la OCUPACIÓN cierra ese caso: cuando se cumplió el nivel de servicio y el "
            "modelo igual pide más gente, es porque en esos intervalos no manda el nivel "
            "de servicio sino el techo de ocupación —en día hábil fija el 40% del "
            "requerimiento; sábado y domingo, cero—. Las dos ocupaciones usan la misma "
            "definición de Erlang (tráfico ofrecido sobre operadores en la cola) y por eso "
            "se pueden comparar; NO es la ocupación del tablero de la operación, que mide "
            "otra cosa y da 84% incluso a las 3 de la mañana. "
            "Se dimensiona con la configuración de HOY, a propósito."))

    if dg.get("motivos"):
        fila = _bloque(ws, fila, "Qué restricción fijó la dotación",
                       [(f"  {m['motivo']}", m["parte"], PCT)
                        for m in dg["motivos"]], nota=(
            "Cada intervalo se dimensiona contra cuatro restricciones a la vez —nivel de "
            "servicio, techo de ocupación, techo de abandono de cada skill y cobertura "
            "mínima del pool— y queda el número de la MÁS EXIGENTE: ésa es la que mandó. "
            "El peso va en intervalos-operador y no en cantidad de intervalos, porque los "
            "intervalos donde manda el techo son los del pico: medido el 07/09, el techo "
            "mandó en 16 intervalos contra 26 del nivel de servicio, pero pesó 985 "
            "intervalos-operador contra 285. El techo de ocupación NO es un compromiso "
            "con el cliente sino una decisión de condiciones de trabajo, así que es la "
            "perilla a mirar si hay que bajar dotación; el nivel de servicio y el "
            "abandono sí lo son."))

    for etiqueta, clave in (("Peor día", "peor_dia"), ("Mejor día", "mejor_dia")):
        d = bt.get(clave)
        if d:
            fila = _bloque(ws, fila, etiqueta, [
                ("Día", d.get("dia"), FECHA),
                ("Real", d.get("real"), NUM0),
                ("Pronosticado", d.get("pronosticado"), NUM0),
                ("Desvío", d.get("desvio"), PCT),
                ("Evento", d.get("evento") or "—"),
            ])

    avisos = bt.get("avisos") or []
    if avisos:
        ws.cell(row=fila, column=1, value="Avisos").font = _TITULO
        fila += 1
        for a in avisos:
            celda = ws.cell(row=fila, column=1, value=f"• {a}")
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=8)
            fila += 1

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 22

    por_tipo = bt.get("por_tipo_de_dia") or []
    if por_tipo:
        _escribir_tabla(wb.create_sheet("Por tipo de día"), _COLS_TIPO_DIA,
                        [_fila_tipo_dia(f) for f in por_tipo])
    _escribir_tabla(wb.create_sheet("Por día"), _COLS_DIA,
                    [_fila_dia(f) for f in bt.get("por_dia") or []])
    dot = bt.get("dotacion") or {}
    if dot.get("por_dia"):
        _escribir_tabla(wb.create_sheet("Operadores"), _con_refuerzo(_COLS_DOTACION, refuerzo),
                        [_fila_dotacion(f) for f in dot["por_dia"]])
    if dot.get("por_intervalo"):
        _escribir_tabla(wb.create_sheet("Operadores por intervalo"),
                        _con_refuerzo(_COLS_DOTACION_INTERVALO, refuerzo),
                        [_fila_dotacion_intervalo(f) for f in dot["por_intervalo"]])
    _escribir_tabla(wb.create_sheet("Por intervalo"), _COLS_INTERVALO,
                    [_fila_intervalo(f) for f in bt.get("por_intervalo") or []])
    return _guardar(wb)


# El cotejo de dotación: cuántos operadores pedíamos y cuántos hubo. Todo en
# intervalos-operador (la suma de los 48 intervalos del día), que es la unidad en
# la que se compran horas.
_COLS_DOTACION = [
    Columna("dia", "Día", FECHA, 12),
    Columna("tipo_dia", "Tipo de día", None, 12),
    Columna("llamadas", "Llamadas reales", NUM0, 15),
    # EN LÍNEA: operadores sobre la cola. Estas se comparan entre sí.
    Columna("en_linea_pron", "Pedíamos en línea (pronóstico)", NUM0, 28),
    Columna("en_linea_real", "Hacía falta en línea (lo que entró)", NUM0, 32),
    Columna("linea_pron_vs_real", "En línea: pedíamos vs hacía falta", PCT, 31),
    Columna("conectados", "Conectados logueados (total)", NUM0, 27),
    Columna("linea_real_vs_conectados", "Hacía falta en línea vs logueados", PCT, 33),
    # En la línea de verdad: logueados menos break y demás pausas.
    Columna("conectados_pool_en_linea", "Telefónicos en línea (sin pausas)", NUM1, 31),
    Columna("conectados_refuerzo_en_linea", "Digital en línea (sin pausas)", NUM1, 28),
    Columna("conectados_digital_en_linea", "Otras digitales en línea (sin pausas)", NUM1, 34),
    Columna("conectados_dedicada_en_linea", "Colas dedicadas en línea (sin pausas)", NUM1, 34),
    Columna("conectados_otras_en_linea", "Sin clasificar en línea (sin pausas)", NUM1, 33),
    Columna("conectados_pool", "Telefónicos logueados", NUM1, 21),
    Columna("conectados_refuerzo", "Digital logueados", NUM1, 17),
    Columna("conectados_digital", "Otras digitales logueados", NUM1, 24),
    Columna("conectados_dedicada", "Colas dedicadas logueados", NUM1, 24),
    Columna("conectados_otras", "Sin clasificar logueados", NUM1, 23),
    # A CITAR: personas, con ausentismo, pausas y break adentro. Estas se comparan
    # entre sí, nunca contra las de arriba.
    Columna("a_planificar_pron", "Pedíamos a citar (pronóstico)", NUM0, 28),
    Columna("a_planificar_real", "Hacía falta a citar (lo que entró)", NUM0, 32),
    Columna("citados", "Tenían turno a citar (registro)", NUM0, 29),
    Columna("citados_parciales",
            "  de Contingencia, SVP y Anfitrión en la línea", NUM0, 42),
    Columna("malla", "Malla a citar (lo publicado)", NUM0, 27),
    Columna("registro_vs_malla", "A citar: turnos de más que la malla", PCT, 32),
    Columna("pron_vs_citados", "A citar: pedíamos vs tenían turno", PCT, 31),
    Columna("real_vs_citados", "A citar: hacía falta vs tenían turno", PCT, 33),
    Columna("de_mas_por_pronostico", "A citar: de más por el pronóstico", NUM0, 31),
    Columna("faltante", "A citar: faltó (por intervalo)", NUM0, 28),
    Columna("refuerzo_turno", "Digital con turno", NUM0, 17),
    Columna("refuerzo_en_linea", "De Digital, en la línea", PCT, 22),
    Columna("cubre_refuerzo", "Lo cubría Digital", NUM0, 17),
    Columna("faltante_neto", "Faltó tras Digital", NUM0, 18),
    Columna("nds_real", "NDS real", PCT, 10),
    Columna("abandono_real", "Abandono real", PCT2, 14),
    Columna("cph_pron", "CPH pronosticado", NUM2, 17),
    Columna("cph_real", "CPH real", NUM2, 10),
    Columna("ocupacion_modelo", "Ocupación del modelo", PCT, 20),
    Columna("ocupacion_real", "Ocupación real", PCT, 15),
    Columna("max_ocupacion", "Techo de ocupación", PCT, 18),
    Columna("peso_techo_ocupacion", "Del pedido, lo fija el techo", PCT, 28),
    Columna("mando_nds", "Mandó: nivel de servicio", PCT, 24),
    Columna("mando_ocupacion", "Mandó: techo de ocupación", PCT, 25),
    Columna("mando_abandono", "Mandó: abandono", PCT, 16),
    Columna("mando_minima", "Mandó: cobertura mínima", PCT, 23),
    Columna("poco_volumen", "Poco volumen (no promedia)", None, 25),
]


def _fila_dotacion(f: dict) -> dict:
    """Una fila por día.

    La brecha se informa partida en dos a propósito: «pedíamos vs citados» es lo
    que se ve en pantalla y mezcla dos errores distintos, y «hacía falta vs
    citados» es el mismo cociente sin el error de pronóstico adentro. La resta de
    los dos dimensionamientos —«de más por el pronóstico»— dice cuál de los dos
    hay que arreglar."""
    malla = f.get("malla") or 0
    # Una columna por familia y no un texto con el reparto: así se puede ordenar y
    # graficar en la planilla, que es para lo que alguien la baja.
    reparto = {m["motivo"]: m["parte"] for m in (f.get("motivos") or ())}
    return {**f,
            "mando_nds": reparto.get("nivel de servicio"),
            "mando_ocupacion": reparto.get("techo de ocupación"),
            "mando_abandono": reparto.get("abandono"),
            "mando_minima": reparto.get("cobertura mínima"),
            "tipo_dia": _NOMBRE_TIPO.get(f.get("tipo_dia"), f.get("tipo_dia")),
            # Horas extras y cambios de último momento: `payroll_futuro` no tiene
            # ninguna hora extra cargada, así que "cuántos hubo" sale del registro.
            "registro_vs_malla": (f["citados"] / malla - 1) if malla else None,
            "poco_volumen": "sí" if f.get("poco_volumen") else ""}


_COLS_TIPO_DIA = [
    Columna("tipo", "Tipo de día", None, 14),
    Columna("n", "Días", NUM0, 8),
    Columna("real", "Llamadas reales", NUM0, 16),
    Columna("modelo", "Error del nuestro", PCT, 18),
    Columna("produccion", "Error del cliente", PCT, 18),
    Columna("combinado", "Error del combinado", PCT, 20),
    Columna("ingenuo", "Error de repetir la semana", PCT, 26),
    Columna("gana", "Quién acierta más", None, 20),
]

_NOMBRE_TIPO = {"habil": "Hábil", "sabado": "Sábado", "domingo": "Domingo",
                "feriado": "Feriado"}


def _fila_tipo_dia(f: dict) -> dict:
    """Una fila por tipo de día. El MAPE y no el WAPE: acá cada día pesa uno.

    La columna «quién acierta más» existe para que la hoja se pueda leer sin
    comparar dos porcentajes a mano, que es exactamente lo que nadie hace."""
    def mape(clave):
        v = f.get(clave)
        return v.get("mape") if v else None

    nuestro, suyo = mape("modelo"), mape("produccion")
    if nuestro is None or suyo is None:
        gana = "—"
    elif abs(nuestro - suyo) < 0.02:
        gana = "empatan"
    else:
        gana = ("el nuestro" if nuestro < suyo else "el del cliente")
    return {
        "tipo": _NOMBRE_TIPO.get(f.get("tipo"), f.get("tipo")),
        "n": f.get("n"), "real": f.get("real"),
        "modelo": nuestro, "produccion": suyo,
        "combinado": mape("combinado"), "ingenuo": mape("ingenuo"),
        "gana": gana,
    }


def _fila_dia(f: dict) -> dict:
    dia = f.get("dia")
    if isinstance(dia, datetime):
        dia = dia.date()
    return {
        **f, "dia": dia,
        "dia_semana": _DIAS[dia.isoweekday() - 1] if isinstance(dia, date) else None,
        "diferencia": (round(f["real"] - f["pronosticado"], 1)
                       if f.get("real") is not None and f.get("pronosticado") is not None
                       else None),
        "parcial": "sí" if f.get("parcial") else "",
    }


# El cotejo de dotación sin sumar: cada fila es media hora. El día agregado dice
# CUÁNTO se erró y esto dice DÓNDE.
_COLS_DOTACION_INTERVALO = [
    Columna("dia", "Día", FECHA, 12),
    Columna("hora", "Intervalo", None, 10),
    Columna("llamadas_pron", "Llamadas pronosticadas", NUM1, 22),
    Columna("llamadas_real", "Llamadas reales", NUM1, 15),
    # EN LÍNEA, igual que la pantalla: se comparan entre sí.
    Columna("en_linea_pron", "Pedíamos en línea (pronóstico)", NUM0, 28),
    Columna("en_linea_real", "Hacía falta en línea (lo que entró)", NUM0, 32),
    Columna("conectados", "Conectados logueados (total)", NUM1, 27),
    Columna("diferencia_linea",
            "En línea: sobró / faltó (telefónicos + digitales en línea − hacía falta)", NUM1, 62),
    Columna("conectados_pool_en_linea", "Telefónicos en línea (sin pausas)", NUM1, 31),
    Columna("conectados_refuerzo_en_linea", "Digital en línea (sin pausas)", NUM1, 28),
    Columna("conectados_digital_en_linea", "Otras digitales en línea (sin pausas)", NUM1, 34),
    Columna("conectados_dedicada_en_linea", "Colas dedicadas en línea (sin pausas)", NUM1, 34),
    Columna("conectados_otras_en_linea", "Sin clasificar en línea (sin pausas)", NUM1, 33),
    Columna("conectados_pool", "Telefónicos logueados", NUM1, 21),
    Columna("conectados_refuerzo", "Digital logueados", NUM1, 17),
    Columna("conectados_digital", "Otras digitales logueados", NUM1, 24),
    Columna("conectados_dedicada", "Colas dedicadas logueados", NUM1, 24),
    Columna("conectados_otras", "Sin clasificar logueados", NUM1, 23),
    # A CITAR: personas, con la cadena de descuentos adentro.
    Columna("a_planificar_pron", "Pedíamos a citar (pronóstico)", NUM0, 28),
    Columna("a_planificar_real", "Hacía falta a citar (lo que entró)", NUM0, 32),
    Columna("a_planificar_bruto", "Hacía falta a citar, sin redondear", NUM2, 32),
    Columna("citados", "Tenían turno a citar (registro)", NUM0, 29),
    Columna("citados_parciales",
            "  de Contingencia, SVP y Anfitrión en la línea", NUM0, 42),
    Columna("malla", "Malla a citar (lo publicado)", NUM0, 27),
    Columna("diferencia", "A citar: sobró / faltó (turno − hacía falta)", NUM0, 40),
    Columna("refuerzo_turno", "Digital con turno", NUM0, 17),
    Columna("cubre_refuerzo", "Lo cubría Digital", NUM0, 17),
    Columna("faltante_neto", "Faltó tras Digital", NUM0, 18),
    Columna("entrantes", "Entrantes", NUM0, 11),
    Columna("nds_real", "NDS real", PCT, 10),
    Columna("abandono_real", "Abandono real", PCT, 14),
    Columna("atendidas", "Atendidas", NUM0, 11),
    Columna("cph_pron", "CPH pronosticado", NUM2, 17),
    Columna("cph_real", "CPH real", NUM2, 10),
    Columna("ocupacion_modelo", "Ocupación del modelo", PCT, 20),
    Columna("ocupacion_real", "Ocupación real", PCT, 15),
    Columna("disponibilidad", "Disponibilidad", PCT, 15),
    Columna("shrinkage", "Shrinkage del día", PCT, 17),
    Columna("factor_break", "Break", PCT, 8),
    Columna("motivo", "Qué mandó", None, 20),
    Columna("motivo_detalle", "Qué mandó (detalle)", None, 32),
]


def _fila_dotacion_intervalo(f: dict) -> dict:
    """Una media hora. «Sobró / faltó» va calculado para que la planilla se pueda
    filtrar por los intervalos donde faltó gente sin armar la fórmula a mano."""
    salida = _con_fecha({**f, "Intervalo": f.get("momento")})
    salida["diferencia"] = (f.get("citados") or 0) - (f.get("a_planificar_real") or 0)
    # La misma pregunta en línea: cada resta, entre dos columnas de la misma unidad.
    # Como la pantalla: telefónicos + Digital sin pausas; sin ese dato, los logueados.
    if f.get("conectados_pool_en_linea") is not None:
        en_linea = (f["conectados_pool_en_linea"] + (f.get("conectados_refuerzo_en_linea") or 0)
                    + (f.get("conectados_digital_en_linea") or 0))
    else:
        en_linea = f.get("conectados") or 0
    salida["diferencia_linea"] = round(en_linea - (f.get("en_linea_real") or 0), 1)
    return salida


def _fila_intervalo(f: dict) -> dict:
    salida = _con_fecha({**f, "Intervalo": f.get("momento")})
    if f.get("real") is not None and f.get("pronosticado") is not None:
        salida["diferencia"] = round(f["real"] - f["pronosticado"], 1)
    return salida


def nombre_de_archivo(prefijo: str, desde=None, hasta=None) -> str:
    """Nombre estable y ordenable. Sin espacios: viaja por mail y por Windows."""
    partes = [prefijo]
    for d in (desde, hasta):
        if isinstance(d, (date, datetime)):
            partes.append(d.strftime("%Y%m%d"))
    return "_".join(partes) + ".xlsx"


# =========================================================================
# LIBRO DE REFUERZOS — el pedido que se le manda a RRHH
# =========================================================================

_ACCION = {
    "malla": "Mover la malla (sin costo extra)",
    "horas_extra": "Horas extra programadas",
    "convocatoria": "Convocatoria del mismo día",
}

_ESTADO_TEXTO = {
    "pendiente": "Pendiente",
    "pedido": "Pedido",
    "cubierto": "Cubierto",
    "cubierto_parcial": "Cubierto parcial",
    "no_cubierto": "No cubierto",
    "descartado": "Descartado",
}

_COLS_REFUERZO = [
    Columna("dia", "Día", FECHA, 12),
    Columna("dia_semana", "Día de semana", None, 14),
    Columna("hora_desde", "Desde", None, 9),
    Columna("hora_hasta", "Hasta", None, 9),
    Columna("pool", "Pool", None, 26),
    Columna("faltante_pico", "Faltan (pico)", NUM0, 13),
    Columna("faltante_promedio", "Faltan (promedio)", NUM1, 17),
    Columna("horas_operador", "Horas-operador", NUM1, 15),
    Columna("faltante_pico_alto", "Pico si viene alto", NUM0, 18),
    Columna("horas_operador_alto", "Horas si viene alto", NUM1, 19),
    Columna("antelacion_dias", "Días de antelación", NUM0, 18),
    Columna("accion_texto", "Cómo se resuelve", None, 32),
    Columna("sugerencia", "Sugerencia", None, 38),
    Columna("estado_texto", "Estado", None, 16),
    Columna("horas_cubiertas", "Horas cubiertas", NUM1, 15),
    Columna("llamadas", "Llamadas del tramo", NUM0, 18),
]


def libro_de_refuerzos(datos: Dict[str, Any], campana: str = "") -> io.BytesIO:
    """Los bloques de refuerzo, listos para pedir."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    r = datos.get("resumen") or {}
    por_accion = r.get("por_accion") or {}

    fila = _bloque(ws, 1, f"Refuerzos necesarios{' — ' + campana if campana else ''}", [
        ("Calculado el", datos.get("hoy"), FECHA),
        ("Período del plan", datos.get("desde"), FECHA),
        ("Hasta", datos.get("hasta"), FECHA),
        ("Corrida", datos.get("corrida_id"), NUM0),
        ("Bloques con faltante", r.get("tramos"), NUM0),
        ("Días afectados", r.get("dias"), NUM0),
        ("Horas-operador que faltan", r.get("horas_operador"), NUM1),
        ("Horas-operador si el volumen viene alto", r.get("horas_operador_alto"), NUM1),
        ("Faltante máximo en un intervalo", r.get("faltante_pico"), NUM0),
    ], nota=(
        "Cada bloque es un pedido: un tramo contiguo de un mismo pool y un mismo "
        "día. Se cortan cuando hay un hueco, porque un faltante de 10 a 12 y otro "
        "de 17 a 19 son dos pedidos distintos y unirlos pediría gente para el "
        "mediodía que no hace falta."))

    if por_accion:
        fila = _bloque(ws, fila, "Por cómo se resuelve", [
            (_ACCION.get(k, k), v.get("horas_operador"), NUM1)
            for k, v in sorted(por_accion.items())
        ], nota=(
            "Con cuatro días o más todavía se puede mover la malla, que no cuesta "
            "plata. Entre uno y tres, hora extra programada. El mismo día, "
            "convocatoria: es lo más caro y lo peor para el clima interno, así que "
            "todo lo que aparezca ahí es plata que se pudo haber ahorrado mirando "
            "esto antes."))

    err = datos.get("error_por_horizonte") or {}
    if err:
        fila = _bloque(ws, fila, "Error típico del pronóstico", [
            (f"a {k} día(s) de antelación", v, PCT) for k, v in sorted(err.items())
        ], nota=(
            "Es lo que define el escenario alto. Quedarse corto y quedarse largo no "
            "cuestan lo mismo: largo son horas pagadas de más, corto es nivel de "
            "servicio, abandono y —en Electrodependientes— un compromiso "
            "contractual. Por eso van las dos columnas y la decisión de cuál mirar "
            "es de negocio, no del modelo."))

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22

    _escribir_tabla(wb.create_sheet("Refuerzos"), _COLS_REFUERZO,
                    [_fila_refuerzo(f) for f in datos.get("refuerzos") or []])
    return _guardar(wb)


def _fila_refuerzo(f: dict) -> dict:
    dia = f.get("dia")
    if isinstance(dia, datetime):
        dia = dia.date()
    ped = f.get("pedido") or {}
    estado = ped.get("estado") if ped else "pendiente"
    sug = f.get("turnos_sugeridos") or {}
    sug_texto = sug.get("texto") if isinstance(sug, dict) else (str(sug) if sug else "")
    return {
        **f, "dia": dia,
        "dia_semana": _DIAS[dia.isoweekday() - 1] if isinstance(dia, date) else None,
        "hora_desde": _hora(f.get("desde")),
        "hora_hasta": _hora(f.get("hasta")),
        "accion_texto": _ACCION.get(f.get("accion"), f.get("accion")),
        "sugerencia": sug_texto,
        "estado_texto": _ESTADO_TEXTO.get(estado, estado),
        "horas_cubiertas": ped.get("horas_cubiertas") if ped else None,
    }


def _hora(v) -> Optional[str]:
    return v.strftime("%H:%M") if isinstance(v, datetime) else None
